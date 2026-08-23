#!/usr/bin/env python3
import json
import re
import subprocess
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PDF = Path("/tmp/ct-docx-render/сборник цэ цт англ 2026(1).pdf")
DOCX = Path("/workspace/scratch/f132a6013268/upload/сборник цэ цт англ 2026(1).docx")


def clean(text):
    replacements = {
        "\u00ad": "", "¬": "", "’ s": "’s", "’ ve": "’ve", "’ m": "’m", "’ d": "’d",
        "’ t": "’t", "i/": "is", "/tart": "start", "deveLoped": "developed",
        "popuLar": "popular", "exceptionaL": "exceptional", "Luggage": "luggage",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s*\n\s*", "\n", text)
    return text.strip()


def normalize_labels(text):
    text = re.sub(r"(?m)^\s*A\s+((?:1\).*)?)\n\s*(1[0-2])\s*\.", lambda m: f"A{m.group(2)}. {m.group(1)}", text)
    text = re.sub(r"\bA[lI]\s*(?=[.)\s])", "A1", text)
    text = re.sub(r"\bA[lI]{2}\s*(?=[.)\s])", "A11", text)
    text = re.sub(r"\bA[lI]\s*([4-8])\b", r"A1\1", text)
    text = re.sub(r"\bА\s*([0-9]+)", r"A\1", text)
    text = re.sub(r"\bВ\s*([0-9]+)", r"B\1", text)
    text = re.sub(r"\bA\s+([0-9]+)", r"A\1", text)
    text = re.sub(r"\bB\s+([0-9]+)", r"B\1", text)
    text = text.replace("Аб.", "A6.").replace("Вб.", "B6.")
    text = re.sub(r"(?m)^\s*A\s*([0-9])\s+([0-9])\s*\.", lambda m: f"A{m.group(1)}{m.group(2)}.", text)
    text = re.sub(r"(?m)^\s*A1\s+([0-9])\s*\.", lambda m: f"A1{m.group(1)}.", text)
    text = re.sub(r"(?m)^\s*A1\s+([3-8])(?=\s)", lambda m: f"A1{m.group(1)}.", text)
    text = re.sub(r"(?m)^\s*A\s*[lI1]\s*([3-8])\s*[r.]", lambda m: f"A1{m.group(1)}.", text)
    text = re.sub(r"(?m)^\s*[АA]\s*1\s*([5-8])\s*[r.]", lambda m: f"A1{m.group(1)}.", text)
    text = re.sub(r"(?m)^\s*A\s+1\s+1\)\s*", "A11. 1) ", text)
    text = re.sub(r"(?<=[A-Za-zА-Яа-я’'])\s*([2-5])\s*\)", r" \1)", text)
    text = re.sub(r"(1\)\s*[^\n]*?)\s+\)\s+([^\n]*?\s+3\))", r"\1 2) \2", text)
    text = re.sub(r"(?m)^\s*A(1[3-8])(?=\s)", r"A\1.", text)
    return text


def split_variants(text):
    heading = re.compile(r"(?im)^\s*(?:(ЦТ|цт)\s+)?ВАРИАНТ\s+([1-5])\s*$|^\s*Вариант\s+([1-5])\s*$")
    matches = list(heading.finditer(text))
    chunks = []
    for index, match in enumerate(matches[:10]):
        end = matches[index + 1].start() if index + 1 < len(matches) else text.find("ОТВЕТЫ", match.end())
        end = end if end > match.end() else len(text)
        number = int(match.group(2) or match.group(3))
        collection = "ЦЭ 2026" if index < 5 else "ЦТ 2026"
        chunks.append((collection, number, normalize_labels(text[match.end():end].replace("\f", "\n"))))
    return chunks


def marker(section, number):
    return re.compile(rf"(?m)^\s*{section}\s*{number}\s*(?:[.\-]|(?=\s))")


def task_span(text, section, number, maximum):
    match = marker(section, number).search(text)
    if not match:
        return ""
    end = len(text)
    for next_number in range(number + 1, maximum + 1):
        next_match = marker(section, next_number).search(text, match.end())
        if next_match:
            end = next_match.start()
            break
    return text[match.start():end]


def split_columns(line):
    return [part.strip() for part in re.split(r"\s{2,}", line.strip()) if part.strip()]


def parse_options(segment, count=5):
    lines = [line for line in segment.splitlines() if line.strip()]
    rows = []
    for line in lines:
        chunks = split_columns(line)
        if chunks:
            rows.append(chunks)
    values = {}
    for row in rows:
        targets = list(range(1, len(row) + 1))
        if len(row) == 3:
            targets = [1, 3, 5]
        elif len(row) == 2:
            targets = [2, 4]
        for target, chunk in zip(targets, row):
            chunk = re.sub(r"^A\s*\d+\s*\.\s*", "", chunk)
            label = re.match(r"^([1-5])\s*[.)]\s*(.*)$", chunk)
            if label:
                target, chunk = int(label.group(1)), label.group(2)
            chunk = re.sub(r"^1\)\s*", "", chunk)
            if chunk:
                values.setdefault(target, clean(chunk))
    if all(i in values for i in range(1, count + 1)):
        return [values[i] for i in range(1, count + 1)]
    # Standard inline lists, including Word's rendered "1." markers.
    inline = list(re.finditer(r"(?<![\w(])([1-5])\s*[.)]\s*", segment))
    direct = {}
    for index, match in enumerate(inline):
        end = inline[index + 1].start() if index + 1 < len(inline) else len(segment)
        value = clean(segment[match.end():end])
        value = re.sub(r"\n\s*", " ", value).strip(" .")
        direct.setdefault(int(match.group(1)), value)
    for line in lines:
        line = re.sub(r"^A\s*\d+\s*\.\s*", "", line.strip())
        line_markers = list(re.finditer(r"(?<![\w(])([2-5])\s*[.)]\s*", line))
        if line_markers:
            prefix = clean(line[:line_markers[0].start()])
            first_number = int(line_markers[0].group(1))
            inferred = 1 if first_number in (2, 3) else 2 if first_number == 4 else None
            if prefix and inferred:
                direct.setdefault(inferred, prefix)
    if all(i in direct for i in range(1, count + 1)):
        return [direct[i] for i in range(1, count + 1)]
    return []


def after_instruction(prefix, pattern):
    matches = list(re.finditer(pattern, prefix, re.I | re.S))
    return clean(prefix[matches[-1].end():]) if matches else clean(prefix)


def strip_label(segment, section, number):
    return clean(marker(section, number).sub("", segment, count=1))


def build_questions(chunk):
    part_b_match = re.search(r"(?im)^\s*ЧАСТЬ\s*[BВ]\s*$", chunk)
    part_a = chunk[:part_b_match.start()] if part_b_match else chunk
    part_b = chunk[part_b_match.end():] if part_b_match else ""
    part_a = re.sub(r"(?im)^\s*ЧАСТЬ\s*A\s*$", "", part_a)
    positions = {i: marker("A", i).search(part_a) for i in range(1, 19)}
    positions = {i: m.start() for i, m in positions.items() if m}
    group1 = after_instruction(part_a[:positions.get(1, len(part_a))], r"Прочитайте\s+текст\.\s*Выберите.*?ответа\.")
    group9 = after_instruction(part_a[:positions.get(9, len(part_a))], r"Прочитайте\s+текст\.\s*Выберите.*?ответа\.")
    group15 = after_instruction(part_a[:positions.get(15, len(part_a))], r"Прочитайте\s+текст\s+и\s+выполните.*?\.")
    choose = list(re.finditer(r"Выберите\s+один\s+из\s+предложенных.*?текста\.", group15, re.I | re.S))
    if choose:
        group15 = clean(group15[:choose[-1].start()])
    result = []
    for number in range(1, 19):
        segment = task_span(part_a, "A", number, 18)
        segment = re.split(r"(?im)\n\s*Прочитайте\s+", segment, maxsplit=1)[0]
        count = 4 if number >= 15 else 5
        options = parse_options(segment, count)
        first_option = re.search(r"(?<![\w(])1\s*[.)]\s*", segment)
        task_text = strip_label(segment[:first_option.start()] if first_option else segment, "A", number)
        if number <= 4:
            prompt, context = f"Выберите правильный вариант для A{number}.", group1
        elif number <= 8:
            fragments = all(f"({i})" in task_text for i in range(1, 6))
            prompt = "Укажите номер ошибочного фрагмента." if fragments else "Выберите правильный вариант."
            context = task_text
            if fragments:
                options = ["1", "2", "3", "4", "5"]
        elif number <= 12:
            prompt, context = f"Выберите правильный вариант для A{number}.", group9
        elif number <= 14:
            prompt, context = "Выберите правильный вариант.", task_text
        else:
            prompt, context = task_text, group15
        result.append({"id": f"A{number}", "prompt": prompt, "context": context, "options": options})
    for number in range(1, 19):
        segment = strip_label(task_span(part_b, "B", number, 18), "B", number)
        segment = re.split(r"(?im)^\s*Прочитайте\s+", segment, maxsplit=1)[0]
        result.append({"id": f"B{number}", "context": segment})
    b19_match = re.search(r"Прочитайте\s+текст\.\s*От\s+приведенных.*?изменить\.\s*", part_b, re.I | re.S)
    b9_match = marker("B", 9).search(part_b)
    if b9_match:
        b19_task = marker("B", 19).search(part_b)
        next_instruction = re.search(r"(?im)^\s*Прочитайте\s+текст", part_b[b9_match.end():])
        next_instruction_pos = b9_match.end() + next_instruction.start() if next_instruction else None
        ends = [match.start() for match in (b19_match, b19_task) if match]
        if next_instruction_pos is not None:
            ends.append(next_instruction_pos)
        end = min(ends) if ends else len(part_b)
        common_b9 = clean(part_b[b9_match.start():end])
        common_b9 = re.sub(r"\s*\n\s*", " ", common_b9).strip()
        for question in result:
            if question["id"].startswith("B") and 9 <= int(question["id"][1:]) <= 18:
                question["context"] = common_b9
    common = clean(part_b[b19_match.end():]) if b19_match else ""
    for number in range(19, 23):
        result.append({"id": f"B{number}", "context": common})
    return result


def docx_option_maps():
    """Read answer columns from Word's actual paragraph/table structure."""
    document = Document(DOCX)
    paragraph_objects = list(document.paragraphs)
    paragraphs = [p.text.strip() for p in paragraph_objects]
    starts = [i for i, value in enumerate(paragraphs) if re.fullmatch(r"(?i)(?:ЦТ\s+)?ВАРИАНТ\s*[1-5]", value)]
    maps = []
    for variant_index, start in enumerate(starts[:10]):
        end = starts[variant_index + 1] if variant_index + 1 < len(starts) else len(paragraphs)
        # Keep tabs intact: they encode Word's option columns.
        block = [value.replace("Аб.", "A6.").replace("А6.", "A6.") for value in paragraphs[start:end]]
        first = next((i for i, value in enumerate(block) if re.match(r"^[AА]\s*[lI1]\s*[.]", value)), None)
        stop = next((i for i, value in enumerate(block[first + 1:], first + 1) if re.match(r"(?i)^Прочитайте\s+предложения", value)), None) if first is not None else None
        option_map = {}
        # Shared reading passages are stored as native Word paragraphs. Reading
        # them here prevents preceding parts/instructions from leaking into a
        # later task when the visible PDF contains unusual spacing.
        a1_context = next((clean(value) for value in block if re.search(r"\([AА]\s*[lI1]\s*\)", value) and len(value) > 120), "")
        a9_context = next((clean(value) for value in block if re.search(r"\([AА]\s*9\s*\)", value) and len(value) > 120), "")
        if a1_context:
            for task in range(1, 5):
                option_map[f"A{task}"] = {"context": a1_context}
        if a9_context:
            for task in range(9, 13):
                option_map[f"A{task}"] = {"context": a9_context}
        b9_index = next((i for i, value in enumerate(block) if re.match(r"^\s*[BВ]\s*9\s*[.]", value)), None)
        b18_index = next((i for i, value in enumerate(block) if re.match(r"^\s*[BВ]\s*1\s*8\s*[.]", value)), None)
        if b9_index is not None and b18_index is not None and b18_index >= b9_index:
            b9_lines = [clean(value) for value in block[b9_index:b18_index + 1] if clean(value)]
            if len(b9_lines) == 10:
                normalized_lines = []
                for offset, value in enumerate(b9_lines):
                    value = re.sub(r"^\s*[BbВв]\s*(?:[lI1]\s*)?[0-9OО]\s*[.]?\s*", "", value)
                    normalized_lines.append(f"B{9 + offset}. {value}")
                common_b9 = "\n".join(normalized_lines)
                for task in range(9, 19):
                    option_map[f"B{task}"] = {"context": common_b9}
        b19_context = next((clean(value) for value in block
                            if re.search(r"\(\s*[BВ]\s*1\s*9\s*\)", value)
                            and re.search(r"\(\s*[BВ]\s*2\s*2\s*\)", value)
                            and len(value) > 150), "")
        if b19_context:
            b19_context = re.sub(r"\(\s*[BВ]\s*(1\s*9|2\s*[0-2])\s*\)", lambda match: f"(B{re.sub(r'\s+', '', match.group(1))})", b19_context)
            for task in range(19, 23):
                option_map[f"B{task}"] = {"context": b19_context}
        if first is not None and stop is not None:
            raw = []
            for value in block[first:stop]:
                if not value:
                    continue
                value = re.sub(r"^[AА]\s*(?:[lI1]|[2-4])\s*[.]\s*", "", value)
                value = re.sub(r"^[1-5]\s*[.)]\s*", "", value)
                raw.append(clean(value))
            if len(raw) == 20:
                for task in range(4):
                    option_map.setdefault(f"A{task + 1}", {}).update({"options": [raw[task * 2], raw[task * 2 + 1], raw[8 + task * 2], raw[9 + task * 2], raw[16 + task]]})
        for task in range(5, 9):
            found = next((i for i, value in enumerate(block) if re.match(rf"^[AА]\s*{task}\s*[.\-]", value)), None)
            if found is None:
                continue
            next_found = next((i for i in range(found + 1, len(block)) if re.match(r"^[AА]\s*[6-9]\s*[.\-]|(?i:^Прочитайте\s+текст)", block[i])), len(block))
            cells = []
            for value in block[found + 1:next_found]:
                cells.extend([clean(cell) for cell in value.split("\t") if clean(cell)])
            values = {}
            option_rows = [[clean(cell) for cell in row.split("\t") if clean(cell)] for row in block[found + 1:next_found] if row.strip()]
            pending = None
            for cell in cells:
                lone = re.fullmatch(r"([1-5])\s*[.)]", cell)
                labeled = re.match(r"^([1-5])\s*[.)]\s*(.+)$", cell)
                if lone:
                    pending = int(lone.group(1))
                elif labeled:
                    values[int(labeled.group(1))] = clean(labeled.group(2))
                    pending = None
                elif pending is not None:
                    values[pending] = cell
                    pending = None
                elif 1 not in values:
                    values[1] = cell
            if len(values) == 5:
                pass
            elif len(option_rows) == 1 and len(cells) == 5:
                values = {}
                for position, cell in enumerate(cells, 1):
                    values[position] = re.sub(r"^[1-5]\s*[.)]\s*", "", cell)
            else:
                values = {}
                for row_cells in option_rows:
                    targets = [1, 3, 5] if len(row_cells) == 3 else [2, 4] if len(row_cells) == 2 else []
                    for target, cell in zip(targets, row_cells):
                        values[target] = re.sub(r"^[1-5]\s*[.)]\s*", "", cell)
            if len(values) == 5:
                context = re.sub(rf"^[AА]\s*{task}\s*[.\-]\s*", "", block[found])
                option_map[f"A{task}"] = {"options": [values[i] for i in range(1, 6)], "context": clean(context)}
        absolute_start, absolute_end = start, end
        found14 = next((i for i in range(absolute_start, absolute_end) if re.match(r"^[AА]\s*[lI1]\s*4\s*[.]", paragraphs[i])), None)
        if found14 is not None:
            dialogue = []
            cursor = found14 + 1
            while cursor < absolute_end and paragraph_objects[cursor].style.name == "List Paragraph" and len(dialogue) < 6:
                dialogue.append(clean(paragraphs[cursor]))
                cursor += 1
            choices = [clean(cell) for cell in paragraphs[cursor].split("\t") if clean(cell)] if cursor < absolute_end else []
            choices = [re.sub(r"^[1-5]\s*[.)]\s*", "", cell) for cell in choices]
            if len(dialogue) == 6 and len(choices) == 5:
                context = "\n".join(f"{chr(65 + i)}. {re.sub(r'^[A-FА-Е]\s*[.]\s*', '', value)}" for i, value in enumerate(dialogue))
                option_map["A14"] = {"options": choices, "context": context, "prompt": "Расположите реплики так, чтобы получился связный диалог."}
        found15 = next((i for i in range(absolute_start, absolute_end) if re.match(r"^[AА]\s*[lI1]\s*5\s*[.]", paragraphs[i])), None)
        a15_context = ""
        if found14 is not None and found15 is not None:
            option_line = next((i for i in range(found14 + 1, found15) if re.search(r"(?:^|\s)[12]\s*\)", paragraphs[i]) and re.search(r"(?:^|\s)5\s*\)", paragraphs[i])), found14)
            option_line += 1
            passage = []
            for value in paragraphs[option_line:found15]:
                value = clean(value)
                if not value:
                    continue
                latin = len(re.findall(r"[A-Za-z]", value))
                cyrillic = len(re.findall(r"[А-Яа-яЁё]", value))
                if latin > cyrillic:
                    passage.append(value)
            a15_context = "\n".join(passage)
        for task in range(15, 19):
            found = next((i for i in range(absolute_start, absolute_end) if re.match(rf"^[AА]\s*[lI1]\s*{task - 10}\s*[.]", paragraphs[i])), None)
            if found is None:
                continue
            choices = []
            for p in paragraph_objects[found + 1:absolute_end]:
                if p.style.name == "List Paragraph" and p.text.strip():
                    choices.append(clean(p.text))
                    if len(choices) == 4:
                        break
                elif choices:
                    break
            if len(choices) == 4:
                prompt = re.sub(r"^[AА]\s*[lI1]\s*[5-8]\s*[.]\s*", "", paragraphs[found])
                option_map[f"A{task}"] = {"options": choices, "prompt": clean(prompt), "context": a15_context}
        if a15_context:
            for task in range(15, 19):
                option_map.setdefault(f"A{task}", {})["context"] = a15_context
        maps.append(option_map)
    for variant_index, table in enumerate(document.tables[:10]):
        for row_index, row in enumerate(table.rows[:4]):
            cells = [clean(cell.text) for cell in row.cells]
            values = []
            for cell in cells:
                match = re.search(r"(?:^|\s)([1-5])\s*[.)]\s*(.+)$", cell)
                if match:
                    values.append(clean(match.group(2)))
            if len(values) == 5:
                    maps[variant_index].setdefault(f"A{9 + row_index}", {}).update({"options": values})
    return maps


def main():
    raw = subprocess.check_output(["pdftotext", "-layout", str(PDF), "-"], text=True, errors="replace")
    tests = []
    option_maps = docx_option_maps()
    for variant_index, (collection, number, chunk) in enumerate(split_variants(raw)):
        questions = build_questions(chunk)
        for question in questions:
            if question["id"] in option_maps[variant_index]:
                question.update(option_maps[variant_index][question["id"]])
        a15_context = next((q.get("context", "") for q in questions if q["id"] in {"A15", "A16", "A17", "A18"} and q.get("context")), "")
        if a15_context:
            for question in questions:
                if question["id"] in {"A15", "A16", "A17", "A18"}:
                    question["context"] = a15_context
        tests.append({"year": collection, "number": number, "questions": questions})
    output = ROOT / "data" / "docx-display-tests.js"
    output.write_text("window.DOCX_DISPLAY_TESTS = " + json.dumps(tests, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    print(json.dumps([{"year": t["year"], "number": t["number"], "questions": len(t["questions"]), "missingOptions": [q["id"] for q in t["questions"] if q["id"].startswith("A") and not q.get("options")]} for t in tests], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
